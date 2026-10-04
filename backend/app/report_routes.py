import io, csv, json
from collections import Counter, defaultdict
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import Response
from sqlalchemy import select
from .db import get_db
from .models import *
from .security import current_user, require
from .domain import (
    scoped_clients,
    scoped_proposals,
    proposal_access,
    serial,
    completeness,
)
from .sla import sla_view, business_days
from .audit import record, verify_chain

router = APIRouter(prefix="/api", tags=["Indicadores y reportes"])


def selection(
    db,
    u,
    date_from=None,
    date_to=None,
    client_id="",
    manager_id="",
    analyst_id="",
    product_id="",
    state="",
    currency="",
    branch="",
    activity="",
    profile_id="",
    amount_min=None,
    amount_max=None,
):
    q = scoped_proposals(db, u).join(Client)
    for col, value in [
        (Proposal.client_id, client_id),
        (Client.manager_id, manager_id),
        (Proposal.product_id, product_id),
        (Proposal.state, state),
        (Proposal.currency, currency),
        (Client.branch, branch),
        (Client.activity, activity),
        (Client.profile_id, profile_id),
    ]:
        if value:
            q = q.where(col == value)
    if analyst_id:
        q = q.where(
            Proposal.id.in_(
                select(AnalystAssignment.proposal_id).where(
                    AnalystAssignment.user_id == analyst_id
                )
            )
        )
    if date_from:
        q = q.where(Proposal.created_at >= date_from)
    if date_to:
        q = q.where(Proposal.created_at <= date_to)
    if amount_min is not None:
        q = q.where(Proposal.amount >= amount_min)
    if amount_max is not None:
        q = q.where(Proposal.amount <= amount_max)
    return list(db.scalars(q.order_by(Proposal.created_at.desc()).limit(5000)))


def internal(db, u, permission):
    if u.client_id:
        raise HTTPException(403, "Estos indicadores son internos.")
    require(db, u, permission)


def avg(values):
    return round(sum(values) / len(values), 2) if values else None


@router.get("/reports/dashboard")
def dashboard(
    date_from: datetime | None = None,
    date_to: datetime | None = None,
    client_id: str = "",
    manager_id: str = "",
    analyst_id: str = "",
    product_id: str = "",
    state: str = "",
    currency: str = "",
    branch: str = "",
    activity: str = "",
    profile_id: str = "",
    amount_min: float | None = None,
    amount_max: float | None = None,
    u=Depends(current_user),
    db=Depends(get_db),
):
    internal(db, u, "report.view")
    proposals = selection(
        db,
        u,
        date_from,
        date_to,
        client_id,
        manager_id,
        analyst_id,
        product_id,
        state,
        currency,
        branch,
        activity,
        profile_id,
        amount_min,
        amount_max,
    )
    ids = [p.id for p in proposals]
    states = Counter(p.state for p in proposals)
    decided = states["approved"] + states["rejected"]
    hours = defaultdict(list)
    breach = Counter()
    area_counts = Counter()
    overdue = set()
    upcoming = set()
    for s in db.scalars(select(SLARecord).where(SLARecord.proposal_id.in_(ids))):
        v = sla_view(db, s)
        hours[s.area].append(v["consumed"])
        area_counts[s.area] += 1
        breach[s.area] += int(v["breached"])
        if not s.ended_at and not s.paused_at:
            if v["breached"]:
                overdue.add(s.proposal_id)
            elif v["color"] == "yellow":
                upcoming.add(s.proposal_id)
    reports = list(
        db.scalars(
            select(CommercialReport).where(CommercialReport.proposal_id.in_(ids))
        )
    )
    automatic_quality = [
        r.ai_assessment["score"]
        for r in reports
        if isinstance(r.ai_assessment.get("score"), (int, float))
    ]
    quality = []
    for r in reports:
        score = r.human_assessment.get("score")
        if isinstance(score, (int, float)):
            quality.append(score)
    incidents = db.scalar(
        select(__import__("sqlalchemy").func.count())
        .select_from(DocumentValidation)
        .join(Document)
        .join(ProposalRequirement)
        .where(
            ProposalRequirement.proposal_id.in_(ids),
            DocumentValidation.status == "incorrect",
        )
    )
    total_days = []
    holidays = set(db.scalars(select(Holiday.date)))
    trend = Counter()
    by_manager = defaultdict(Counter)
    by_analyst = defaultdict(Counter)
    for p in proposals:
        trend[p.created_at.strftime("%Y-%m")] += 1
        c = db.get(Client, p.client_id)
        by_manager[c.manager_id][p.state] += 1
        total_days.append(
            business_days(
                p.created_at,
                p.updated_at
                if p.state in {"approved", "rejected", "cancelled"}
                else now(),
                holidays,
            )
        )
        for a in db.scalars(
            select(AnalystAssignment).where(
                AnalystAssignment.proposal_id == p.id,
                AnalystAssignment.active.is_(True),
            )
        ):
            by_analyst[a.user_id][p.state] += 1
    incident_ids = set(
        db.scalars(
            select(ProposalRequirement.proposal_id)
            .join(Document)
            .join(DocumentValidation)
            .where(
                ProposalRequirement.proposal_id.in_(ids),
                DocumentValidation.status == "incorrect",
            )
        )
    )
    return dict(
        total_proposals=len(proposals),
        active_proposals=sum(
            v
            for k, v in states.items()
            if k not in {"approved", "rejected", "cancelled"}
        ),
        in_analysis=states["in_analysis"] + states["reanalysis"],
        approval_rate=round(states["approved"] / decided * 100, 2) if decided else 0,
        rejection_rate=round(states["rejected"] / decided * 100, 2) if decided else 0,
        overdue=len(overdue),
        upcoming=len(upcoming),
        by_state=dict(states),
        document_completeness=avg([completeness(db, p)[0] for p in proposals]) or 0,
        document_incidents=incidents,
        approved_with_incidents=sum(
            p.state == "approved" and p.id in incident_ids for p in proposals
        ),
        approved_without_incidents=sum(
            p.state == "approved" and p.id not in incident_ids for p in proposals
        ),
        avg_report_quality=avg(quality),
        avg_automatic_report_quality=avg(automatic_quality),
        report_quality_source="human",
        avg_commercial_days=avg(hours["commercial"]),
        avg_risk_days=avg(hours["risk"]),
        avg_total_days=avg(total_days),
        commercial_breached_pct=round(
            breach["commercial"] / max(area_counts["commercial"], 1) * 100, 2
        ),
        risk_breached_pct=round(breach["risk"] / max(area_counts["risk"], 1) * 100, 2),
        by_manager=[
            {
                "user_id": k,
                "name": db.get(User, k).name,
                "states": dict(v),
                "total": sum(v.values()),
            }
            for k, v in by_manager.items()
        ],
        by_analyst=[
            {
                "user_id": k,
                "name": db.get(User, k).name,
                "states": dict(v),
                "total": sum(v.values()),
            }
            for k, v in by_analyst.items()
        ],
        trend=[{"period": k, "total": v} for k, v in sorted(trend.items())],
    )


def safe_cell(value):
    text = str(value) if value is not None else ""
    return "'" + text if text.startswith(("=", "+", "-", "@", "\t", "\r")) else text


def export_rows(rows, format, title, filename):
    if format == "csv":
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerows([[safe_cell(v) for v in row] for row in rows])
        body = out.getvalue().encode("utf-8-sig")
        mime = "text/csv"
    elif format == "xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte"
        for row in rows:
            ws.append([safe_cell(v) for v in row])
        out = io.BytesIO()
        wb.save(out)
        body = out.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif format == "pdf":
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Table,
            TableStyle,
            Spacer,
        )
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from xml.sax.saxutils import escape

        out = io.BytesIO()
        styles = getSampleStyleSheet()
        doc = SimpleDocTemplate(
            out, pagesize=landscape(A4), leftMargin=24, rightMargin=24
        )
        grid = [
            [Paragraph(escape(str(v or ""))[:2000], styles["Normal"]) for v in row]
            for row in rows
        ]
        table = Table(grid, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dceaf5")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                ]
            )
        )
        doc.build([Paragraph(escape(title), styles["Title"]), Spacer(1, 12), table])
        body = out.getvalue()
        mime = "application/pdf"
    elif format == "docx":
        from docx import Document as WordDocument

        doc = WordDocument()
        doc.add_heading(title, 0)
        for row in rows:
            doc.add_paragraph(" · ".join(str(v or "") for v in row))
        out = io.BytesIO()
        doc.save(out)
        body = out.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        raise HTTPException(422, "Seleccione CSV, Excel XLSX o PDF.")
    return Response(
        body,
        media_type=mime,
        headers={"Content-Disposition": f'attachment; filename="{filename}.{format}"'},
    )


@router.get("/reports/export")
def reports_export(
    request: Request,
    format: str = "csv",
    date_from: datetime | None = None,
    date_to: datetime | None = None,
    client_id: str = "",
    manager_id: str = "",
    analyst_id: str = "",
    product_id: str = "",
    state: str = "",
    currency: str = "",
    branch: str = "",
    activity: str = "",
    profile_id: str = "",
    amount_min: float | None = None,
    amount_max: float | None = None,
    u=Depends(current_user),
    db=Depends(get_db),
):
    internal(db, u, "report.export")
    proposals = selection(
        db,
        u,
        date_from,
        date_to,
        client_id,
        manager_id,
        analyst_id,
        product_id,
        state,
        currency,
        branch,
        activity,
        profile_id,
        amount_min,
        amount_max,
    )
    rows = [
        [
            "Expediente",
            "Empresa",
            "Producto",
            "Importe",
            "Moneda",
            "Estado",
            "Gestor",
            "Días comercial",
            "Días riesgos",
            "SLA comercial",
            "SLA riesgos",
            "Motivo decisión",
        ]
    ]
    for p in proposals:
        c = db.get(Client, p.client_id)
        sla = {}
        for s in db.scalars(
            select(SLARecord)
            .where(SLARecord.proposal_id == p.id)
            .order_by(SLARecord.created_at)
        ):
            sla[s.area] = sla_view(db, s)
        decision = db.scalar(
            select(CreditDecision)
            .where(CreditDecision.proposal_id == p.id)
            .order_by(CreditDecision.created_at.desc())
            .limit(1)
        )
        rows.append(
            [
                p.code,
                c.name,
                db.get(Product, p.product_id).name,
                str(p.amount),
                p.currency,
                p.state,
                db.get(User, c.manager_id).name,
                sla.get("commercial", {}).get("consumed", ""),
                sla.get("risk", {}).get("consumed", ""),
                sla.get("commercial", {}).get("color", ""),
                sla.get("risk", {}).get("color", ""),
                decision.justification if decision else "",
            ]
        )
    result = export_rows(
        rows, format, "WorkFlow Créditos · Operación", "operacion-creditos"
    )
    record(
        db,
        u,
        "report.export",
        "reports",
        after={"format": format, "rows": len(proposals)},
        request=request,
    )
    db.commit()
    return result


@router.get("/proposals/{pid}/report/export")
def report_export(
    pid: str,
    request: Request,
    format: str = "pdf",
    u=Depends(current_user),
    db=Depends(get_db),
):
    internal(db, u, "report.export")
    p = proposal_access(db, u, pid)
    r = db.scalar(select(CommercialReport).where(CommercialReport.proposal_id == pid))
    if not r:
        raise HTTPException(404, "El informe todavía no fue elaborado.")
    result = export_rows(
        [["Informe comercial", "Versión"], [r.content, r.version]],
        format,
        p.code + " · Informe comercial",
        p.code + "-informe",
    )
    record(db, u, "report.download", r.id, after={"format": format}, request=request)
    db.commit()
    return result


@router.get("/audit/export")
def audit_export(
    request: Request, format: str = "csv", u=Depends(current_user), db=Depends(get_db)
):
    internal(db, u, "audit.export")
    events = list(db.scalars(select(AuditEvent).order_by(AuditEvent.id).limit(20000)))
    rows = [
        ["Id", "Fecha", "Usuario", "Acción", "Recurso", "Anterior", "Posterior", "Hash"]
    ] + [
        [
            e.id,
            e.created_at.isoformat(),
            e.actor_id,
            e.action,
            e.resource,
            json.dumps(e.before, ensure_ascii=False),
            json.dumps(e.after, ensure_ascii=False),
            e.hash,
        ]
        for e in events
    ]
    result = export_rows(rows, format, "Auditoría WorkFlow Créditos", "auditoria")
    record(
        db,
        u,
        "audit.export",
        "audit",
        after={"format": format, "events": len(events)},
        request=request,
    )
    db.commit()
    return result


@router.get("/audit/verify")
def audit_verify(u=Depends(current_user), db=Depends(get_db)):
    internal(db, u, "audit.view")
    events = list(db.scalars(select(AuditEvent).order_by(AuditEvent.id)))
    return {"valid": verify_chain(events), "events": len(events)}


@router.get("/audit")
def audit_events(
    action: str = "",
    resource: str = "",
    offset: int = 0,
    u=Depends(current_user),
    db=Depends(get_db),
):
    internal(db, u, "audit.view")
    q = select(AuditEvent)
    if action:
        q = q.where(AuditEvent.action == action)
    if resource:
        q = q.where(AuditEvent.resource == resource)
    return [
        {
            **serial(e),
            "user_name": db.get(User, e.actor_id).name
            if e.actor_id and db.get(User, e.actor_id)
            else "Sistema",
        }
        for e in db.scalars(
            q.order_by(AuditEvent.id.desc()).offset(max(0, offset)).limit(200)
        )
    ]


@router.get("/notifications")
def notifications(u=Depends(current_user), db=Depends(get_db)):
    require(db, u, "notification.view")
    return [
        serial(n)
        for n in db.scalars(
            select(Notification)
            .where(Notification.user_id == u.id)
            .order_by(Notification.created_at.desc())
            .limit(200)
        )
    ]


@router.post("/notifications/{nid}/read")
def read_notification(nid: str, u=Depends(current_user), db=Depends(get_db)):
    require(db, u, "notification.view")
    n = db.get(Notification, nid)
    if not n or n.user_id != u.id:
        raise HTTPException(404, "Notificación no encontrada.")
    n.read = True
    db.commit()
    return {"read": True}
