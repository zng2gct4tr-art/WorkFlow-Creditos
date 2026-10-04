import csv, io
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File, Form
from sqlalchemy import select, or_
from pydantic import ValidationError
from .db import get_db
from .models import *
from .security import current_user, require, permissions
from .schemas import ClientInput
from .domain import (
    client_access,
    scoped_clients,
    scoped_proposals,
    serial,
    proposal_dict,
)
from .audit import record

router = APIRouter(prefix="/api", tags=["Clientes"])


def client_out(db, c, u):
    data = serial(c)
    if u.client_id:
        data = {
            k: v
            for k, v in data.items()
            if k
            in {
                "id",
                "name",
                "rut",
                "account_number",
                "contact",
                "email",
                "phone",
                "address",
                "activity",
                "version",
            }
        }
    else:
        data["manager_name"] = db.get(User, c.manager_id).name
    return data


def validate_client_refs(db, u, body):
    require(db, u, "client.create")
    if u.client_id:
        raise HTTPException(403, "Acción interna.")
    manager = db.get(User, body.manager_id)
    if (
        not manager
        or not manager.active
        or manager.client_id
        or "proposal.create" not in permissions(db, manager)
    ):
        raise HTTPException(422, "Seleccione un gestor comercial activo.")
    if "portfolio.all" not in permissions(db, u) and body.manager_id != u.id:
        raise HTTPException(403, "Solo puede incorporar clientes a su propia cartera.")
    profile = db.get(Profile, body.profile_id)
    if not profile or not profile.active:
        raise HTTPException(422, "Seleccione un perfil documental activo.")


@router.get("/catalog")
def catalog(u=Depends(current_user), db=Depends(get_db)):
    products = [
        serial(p) for p in db.scalars(select(Product).where(Product.active.is_(True)))
    ]
    profiles = [
        serial(p) for p in db.scalars(select(Profile).where(Profile.active.is_(True)))
    ]
    users = []
    if not u.client_id:
        # Only directory needed for assignment; names and identifiers, no credentials.
        perms = permissions(db, u)
        candidates = (
            db.scalars(
                select(User).where(User.active.is_(True), User.client_id.is_(None))
            )
            if ("assignment.manage" in perms or "user.manage" in perms)
            else [u]
        )
        users = [
            {
                "id": v.id,
                "name": v.name,
                "can_analyze": "risk.approve" in permissions(db, v),
            }
            for v in candidates
        ]
    return {"products": products, "profiles": profiles, "users": users}


@router.get("/clients")
def clients(q: str = "", u=Depends(current_user), db=Depends(get_db)):
    query = scoped_clients(db, u)
    if q:
        query = query.where(
            or_(
                Client.name.ilike("%" + q[:200] + "%"),
                Client.rut.ilike("%" + q[:200] + "%"),
                Client.account_number.ilike("%" + q[:200] + "%"),
            )
        )
    return [
        client_out(db, c, u)
        for c in db.scalars(query.order_by(Client.created_at.desc()).limit(500))
    ]


@router.post("/clients", status_code=201)
def create_client(
    body: ClientInput, request: Request, u=Depends(current_user), db=Depends(get_db)
):
    validate_client_refs(db, u, body)
    c = Client(**body.model_dump(mode="json"))
    db.add(c)
    db.flush()
    if body.contact:
        db.add(
            ClientContact(
                client_id=c.id,
                name=body.contact,
                email=str(body.email),
                phone=body.phone,
            )
        )
    record(db, u, "client.create", c.id, after=serial(c), request=request)
    db.commit()
    return client_out(db, c, u)


@router.get("/clients/{cid}")
def detail(cid: str, request: Request, u=Depends(current_user), db=Depends(get_db)):
    c = client_access(db, u, db.get(Client, cid))
    record(db, u, "client.view", c.id, request=request)
    db.commit()
    data = client_out(db, c, u)
    data["proposals"] = [
        proposal_dict(db, p, u)
        for p in db.scalars(
            scoped_proposals(db, u)
            .where(Proposal.client_id == cid)
            .order_by(Proposal.created_at.desc())
        )
    ]
    return data


@router.put("/clients/{cid}")
def edit(
    cid: str,
    body: ClientInput,
    request: Request,
    u=Depends(current_user),
    db=Depends(get_db),
):
    require(db, u, "client.edit")
    c = client_access(
        db, u, db.scalar(select(Client).where(Client.id == cid).with_for_update())
    )
    validate_client_refs(db, u, body)
    before = serial(c)
    for k, v in body.model_dump(mode="json").items():
        setattr(c, k, v)
    c.version += 1
    record(db, u, "client.edit", c.id, before=before, after=serial(c), request=request)
    db.commit()
    return client_out(db, c, u)


@router.get("/import/template")
def import_template(u=Depends(current_user), db=Depends(get_db)):
    from fastapi.responses import Response

    require(db, u, "client.import")
    if u.client_id:
        raise HTTPException(403, "La importación es exclusiva del equipo interno.")
    return Response(
        "name,rut,account_number,manager_id,profile_id,contact,email,phone,address,activity\n",
        media_type="text/csv",
        headers={
            "Content-Disposition": 'attachment; filename="plantilla-clientes.csv"'
        },
    )


@router.post("/import/clients")
async def import_clients(
    file: UploadFile = File(...),
    commit: bool = Form(False),
    request: Request = None,
    u=Depends(current_user),
    db=Depends(get_db),
):
    require(db, u, "client.import")
    if u.client_id:
        raise HTTPException(403, "Acción interna.")
    raw = await file.read(10 * 1024 * 1024 + 1)
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(413, "La importación supera el límite técnico de 10 MB.")
    filename = (file.filename or "").lower()
    try:
        if filename.endswith(".csv"):
            rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
        elif filename.endswith(".xlsx"):
            # Worksheet limits are enforced before parsing to prevent decompression abuse.
            import zipfile

            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if (
                    sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024
                    or len(z.infolist()) > 1000
                ):
                    raise ValueError("Archivo comprimido demasiado grande.")
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator)
            rows = [dict(zip(headers, row)) for _, row in zip(range(10001), iterator)]
            workbook.close()
        else:
            raise HTTPException(422, "Use archivos CSV UTF-8 o Excel XLSX.")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            422, "No se pudo leer el archivo. Revise el formato y las columnas."
        )
    if len(rows) > 10000:
        raise HTTPException(
            413, "La importación supera las 10.000 filas. Divida el archivo."
        )
    results = []
    valid = []
    seen = set()
    for i, row in enumerate(rows, 2):
        try:
            values = {
                k: str(v) if v is not None else ""
                for k, v in row.items()
                if k in ClientInput.model_fields
            }
            values.setdefault("manager_id", u.id)
            body = ClientInput(**values)
            validate_client_refs(db, u, body)
            if body.rut in seen or db.scalar(
                select(Client.id).where(
                    or_(
                        Client.rut == body.rut,
                        Client.account_number == body.account_number,
                    )
                )
            ):
                raise ValueError("RUT o cuenta duplicados.")
            if any(b.account_number == body.account_number for b in valid):
                raise ValueError("Cuenta duplicada en el archivo.")
            seen.add(body.rut)
            valid.append(body)
            results.append(
                {"row": i, "valid": True, "name": body.name, "rut": body.rut}
            )
        except Exception as e:
            message = (
                e.detail
                if isinstance(e, HTTPException)
                else "Revise los campos requeridos, el RUT de 12 dígitos y los duplicados."
                if isinstance(e, ValidationError)
                else str(e)
            )
            results.append({"row": i, "valid": False, "error": message})
    if commit:
        if any(not r["valid"] for r in results):
            raise HTTPException(
                422,
                {
                    "message": "Corrija los errores antes de importar. No se guardó ninguna fila.",
                    "rows": results,
                },
            )
        for body in valid:
            c = Client(**body.model_dump(mode="json"))
            db.add(c)
            db.flush()
            record(db, u, "client.import", c.id, after=serial(c), request=request)
        db.commit()
    return {
        "rows": results,
        "valid": len(valid),
        "errors": len(rows) - len(valid),
        "imported": len(valid) if commit else 0,
    }


@router.get("/search")
def search(q: str, u=Depends(current_user), db=Depends(get_db)):
    ids = list(db.scalars(scoped_clients(db, u).with_only_columns(Client.id)))
    pids = list(db.scalars(scoped_proposals(db, u).with_only_columns(Proposal.id)))
    term = "%" + q[:200] + "%"
    return {
        "clients": [
            client_out(db, c, u)
            for c in db.scalars(
                select(Client)
                .where(
                    Client.id.in_(ids),
                    or_(
                        Client.name.ilike(term),
                        Client.rut.ilike(term),
                        Client.account_number.ilike(term),
                    ),
                )
                .limit(30)
            )
        ],
        "proposals": [
            {
                "id": p.id,
                "code": p.code,
                "client_name": db.get(Client, p.client_id).name,
            }
            for p in db.scalars(
                select(Proposal)
                .where(Proposal.id.in_(pids), Proposal.code.ilike(term))
                .limit(30)
            )
        ],
        "documents": [
            {"id": d.id, "name": d.name, "proposal_id": r.proposal_id}
            for d, r in db.execute(
                select(Document, ProposalRequirement)
                .join(
                    ProposalRequirement,
                    Document.requirement_id == ProposalRequirement.id,
                )
                .join(Proposal, Proposal.id == ProposalRequirement.proposal_id)
                .where(
                    Proposal.id.in_(pids),
                    Document.deleted.is_(False),
                    Document.name.ilike(term),
                    ProposalRequirement.client_visible.is_(True)
                    if u.client_id
                    else __import__("sqlalchemy").true(),
                )
                .limit(30)
            )
        ],
    }
